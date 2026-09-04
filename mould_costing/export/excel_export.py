"""Writes the full line-item costing breakdown to a new Excel workbook."""

from openpyxl import Workbook
from openpyxl.styles import Font


def export_summary(path: str, sections: list, grand_total: float) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Costing Summary"

    all_columns: list[str] = ["Section"]
    for section in sections:
        for col in section.full_columns:
            if col not in all_columns:
                all_columns.append(col)

    ws.append(all_columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for section in sections:
        for row in section.get_export_rows():
            ws.append([row.get(col, "") for col in all_columns])

    ws.append([])
    total_col_index = all_columns.index("Total") if "Total" in all_columns else 1
    total_row = [""] * len(all_columns)
    total_row[0] = "Grand Total"
    total_row[total_col_index] = grand_total
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = max(10, length + 2)

    wb.save(path)
