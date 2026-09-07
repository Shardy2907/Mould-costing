"""Generates an empty standard-chart workbook (headers only) for the user to fill in.

Reads the sheet/column schema directly from chart_loader.REQUIRED_SHEETS so the
template can never drift out of sync with what the app actually expects.

Run with: uv run scripts/make_chart_template.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font

from mould_costing.data.chart_loader import REQUIRED_SHEETS

OUTPUT_PATH = Path(__file__).parent.parent / "Standard_Chart_Template.xlsx"


def build() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, columns in REQUIRED_SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col_index, column_name in enumerate(columns, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_index).column_letter].width = max(
                12, len(column_name) + 2
            )

    wb.save(OUTPUT_PATH)
    print(f"Empty chart template written to {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
