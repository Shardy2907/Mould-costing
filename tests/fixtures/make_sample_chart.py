"""Generates a small sample chart workbook matching the schema in chart_loader.py,
purely for development/testing until the real standardized chart is ready.

Run with: uv run tests/fixtures/make_sample_chart.py
"""

from pathlib import Path

from openpyxl import Workbook

OUTPUT_PATH = Path(__file__).parent / "sample_chart.xlsx"


def build() -> None:
    wb = Workbook()

    guide_pin = wb.active
    guide_pin.title = "GuidePin"
    guide_pin.append(["Type", "Diameter C", "Diameter F", "Diameter G", "Length", "Total Cost"])
    guide_pin.append(["A", 10, 8, 6, 50, 120.0])
    guide_pin.append(["A", 10, 8, 6, 100, 160.0])
    guide_pin.append(["B", 12, 10, 8, 50, 140.0])
    guide_pin.append(["C", 16, 13, 10, 100, 200.0])

    guide_bush = wb.create_sheet("GuideBush")
    guide_bush.append(["Type", "Diameter C", "Diameter D", "Diameter E", "Length", "Total Cost"])
    guide_bush.append(["A", 10, 8, 6, 50, 90.0])
    guide_bush.append(["A", 10, 8, 6, 100, 130.0])
    guide_bush.append(["B", 12, 10, 8, 50, 110.0])

    return_pin = wb.create_sheet("ReturnPin")
    return_pin.append(["D1", "D2", "K", "Length", "Rate"])
    return_pin.append([10, 16, 6, 50, 60.0])
    return_pin.append([12, 18, 7, 63, 75.0])
    return_pin.append([16, 22, 8, 100, 110.0])

    bolts = wb.create_sheet("Bolts")
    bolts.append(["D1", "Length", "Rate"])
    bolts.append(["M6", 20, 5.0])
    bolts.append(["M8", 25, 7.5])
    bolts.append(["M10", 30, 10.0])

    dowelling_sleeve = wb.create_sheet("DowellingSleeve")
    dowelling_sleeve.append(["Outer dia", "Length", "Rate"])
    dowelling_sleeve.append([8, 20, 15.0])
    dowelling_sleeve.append([10, 25, 18.0])

    hook_strip = wb.create_sheet("HookStrip")
    hook_strip.append(["W", "T", "Length", "Rate"])
    hook_strip.append([25, 6, 100, 40.0])
    hook_strip.append([32, 8, 150, 55.0])

    ejector_guide_pin = wb.create_sheet("EjectorGuidePin")
    ejector_guide_pin.append(["D1", "Length", "Rate"])
    ejector_guide_pin.append([8, 50, 35.0])
    ejector_guide_pin.append([10, 63, 42.0])

    ejector_guide_bush = wb.create_sheet("EjectorGuideBush")
    ejector_guide_bush.append(["D1", "Length", "Rate"])
    ejector_guide_bush.append([8, 50, 30.0])
    ejector_guide_bush.append([10, 63, 38.0])

    locating_ring = wb.create_sheet("LocatingRing")
    locating_ring.append(["D1", "thickness", "cost"])
    locating_ring.append([110, 15, 600.0])

    plate_thickness = wb.create_sheet("PlateThickness")
    plate_thickness.append(["Thickness"])
    for value in (16, 20, 25, 32, 40, 50, 63, 80, 100, 125, 150):
        plate_thickness.append([value])

    cavity_housing_drilling = wb.create_sheet("Cavityhousingdrilling")
    cavity_housing_drilling.append(["Diameter", "Thickness", "Cost"])
    for diameter in ("M10", "M12", "M16"):
        cavity_housing_drilling.append([diameter, "24-30", 185])
        cavity_housing_drilling.append([diameter, "38-48", 195])

    spacer_blocks_costing = wb.create_sheet("SpacerBlocksCosting")
    spacer_blocks_costing.append(["Thickness", "Cost"])
    spacer_blocks_costing.append(["24-30", 570])
    spacer_blocks_costing.append(["30-48", 580])
    spacer_blocks_costing.append(["48-60", 590])
    spacer_blocks_costing.append(["60-80", 610])
    spacer_blocks_costing.append(["80-100", 630])

    punch_housing_drilling = wb.create_sheet("PunchHousingDrilling")
    punch_housing_drilling.append(["Diameter", "Thickness", "Cost"])
    for diameter in ("M10", "M12", "M16"):
        punch_housing_drilling.append([diameter, "24-30", 970])
        punch_housing_drilling.append([diameter, "38-48", 1030])
        punch_housing_drilling.append([diameter, "48-60", 1090])

    punch_back_plate = wb.create_sheet("PunchBackPlate")
    punch_back_plate.append(["Diameter", "Thickness", "Cost"])
    for diameter in ("M10", "M12", "M16"):
        punch_back_plate.append([diameter, "24-30", 950])
        punch_back_plate.append([diameter, "38-48", 980])
        punch_back_plate.append([diameter, "48-60", 1010])

    bolt_standards = wb.create_sheet("BoltStandards")
    bolt_standards.append(["Bolt Dia", "Clearance hole dia", "Counter bore dia", "Counter bore depth"])
    bolt_standards.append(["M10", 11, 18, 11])
    bolt_standards.append(["M12", 14, 20, 13])
    bolt_standards.append(["M16", 18, 26, 18])

    wb.save(OUTPUT_PATH)
    print(f"Sample chart written to {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
